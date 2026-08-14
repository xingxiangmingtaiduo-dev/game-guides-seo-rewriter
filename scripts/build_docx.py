from __future__ import annotations

import argparse
import hashlib
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


def metadata_to_dict(metadata):
    return {key: value for key, value in metadata}


def normalize_output_language(value: str) -> str:
    normalized = value.strip().lower()
    aliases = {
        "en": "english",
        "english": "english",
        "es": "spanish",
        "español": "spanish",
        "espanol": "spanish",
        "spanish": "spanish",
        "pt": "portuguese",
        "portuguese": "portuguese",
        "português": "portuguese",
        "portugues": "portuguese",
        "zh": "chinese",
        "zh-cn": "chinese",
        "chinese": "chinese",
        "simplified chinese": "chinese",
        "简体中文": "chinese",
        "簡體中文": "chinese",
        "zh-tw": "traditional_chinese",
        "zh-hk": "traditional_chinese",
        "traditional chinese": "traditional_chinese",
        "繁体中文": "traditional_chinese",
        "繁體中文": "traditional_chinese",
        "th": "thai",
        "thai": "thai",
        "ภาษาไทย": "thai",
        "id": "indonesian",
        "id-id": "indonesian",
        "indonesian": "indonesian",
        "bahasa indonesia": "indonesian",
        "vi": "vietnamese",
        "vi-vn": "vietnamese",
        "vietnamese": "vietnamese",
        "tiếng việt": "vietnamese",
        "tieng viet": "vietnamese",
    }
    return aliases.get(normalized, normalized or "english")


def get_language_pack(output_language: str):
    language = normalize_output_language(output_language)
    packs = {
        "english": {
            "introduction": "Introduction",
            "conclusion": "Conclusion",
            "metadata_heading": "SEO Metadata",
            "metadata_field": "Field",
            "metadata_value": "Value",
            "title_subtitle": "Generated from the SEO article package. Preserve source screenshots inline when available.",
            "source_screenshots": "Source Screenshots",
            "source_caption_prefix": "Source screenshot",
            "image_plan": "Image Plan",
            "image_label": "Image",
            "image_alt": "Alt Text",
            "image_purpose": "Purpose",
            "additional_source_screenshots": "Additional Source Screenshots",
            "appendix_note": "These screenshots were preserved from the source document but not placed in the main article flow.",
            "appendix_images": "Appendix Images",
        },
        "spanish": {
            "introduction": "Introducción",
            "conclusion": "Conclusión",
            "metadata_heading": "Metadatos SEO",
            "metadata_field": "Campo",
            "metadata_value": "Valor",
            "title_subtitle": "Generado a partir del paquete SEO del artículo. Conserva capturas originales cuando están disponibles.",
            "source_screenshots": "Capturas originales",
            "source_caption_prefix": "Captura original",
            "image_plan": "Plan de imágenes",
            "image_label": "Imagen",
            "image_alt": "Texto alt",
            "image_purpose": "Uso",
            "additional_source_screenshots": "Capturas originales adicionales",
            "appendix_note": "Estas capturas se conservaron del documento fuente, pero no se ubicaron en el flujo principal del artículo.",
            "appendix_images": "Imágenes del apéndice",
        },
        "portuguese": {
            "introduction": "Introdução",
            "conclusion": "Conclusão",
            "metadata_heading": "Metadados SEO",
            "metadata_field": "Campo",
            "metadata_value": "Valor",
            "title_subtitle": "Gerado a partir do pacote SEO do artigo. Mantém capturas originais quando disponíveis.",
            "source_screenshots": "Capturas de origem",
            "source_caption_prefix": "Captura de origem",
            "image_plan": "Plano de imagens",
            "image_label": "Imagem",
            "image_alt": "Texto alt",
            "image_purpose": "Uso",
            "additional_source_screenshots": "Capturas de origem adicionais",
            "appendix_note": "Estas capturas foram preservadas do documento de origem, mas não foram colocadas no fluxo principal do artigo.",
            "appendix_images": "Imagens do apêndice",
        },
        "chinese": {
            "introduction": "引言",
            "conclusion": "结语",
            "metadata_heading": "SEO 元数据",
            "metadata_field": "字段",
            "metadata_value": "内容",
            "title_subtitle": "根据 SEO 文章包生成，原文截图已按内容位置保留。",
            "source_screenshots": "原文截图",
            "source_caption_prefix": "原图",
            "image_plan": "图片规划",
            "image_label": "图片",
            "image_alt": "替代文字",
            "image_purpose": "用途",
            "additional_source_screenshots": "其他原文截图",
            "appendix_note": "这些截图来自源文档，但未放入正文流程。",
            "appendix_images": "附录图片",
        },
        "traditional_chinese": {
            "introduction": "前言",
            "conclusion": "結語",
            "metadata_heading": "SEO 資訊",
            "metadata_field": "欄位",
            "metadata_value": "內容",
            "title_subtitle": "根據 SEO 文章包生成，原文截圖已按內容位置保留。",
            "source_screenshots": "原文截圖",
            "source_caption_prefix": "原圖",
            "image_plan": "圖片規劃",
            "image_label": "圖片",
            "image_alt": "替代文字",
            "image_purpose": "用途",
            "additional_source_screenshots": "其他原文截圖",
            "appendix_note": "這些截圖來自源文件，但未放入正文流程。",
            "appendix_images": "附錄圖片",
        },
        "thai": {
            "introduction": "บทนำ",
            "conclusion": "สรุป",
            "metadata_heading": "ข้อมูล SEO",
            "metadata_field": "ฟิลด์",
            "metadata_value": "เนื้อหา",
            "title_subtitle": "สร้างจากแพ็กเกจบทความ SEO และคงภาพต้นฉบับไว้ตามตำแหน่งเนื้อหา",
            "source_screenshots": "ภาพจากต้นฉบับ",
            "source_caption_prefix": "ภาพต้นฉบับ",
            "image_plan": "แผนภาพ",
            "image_label": "ภาพ",
            "image_alt": "ข้อความ Alt",
            "image_purpose": "วัตถุประสงค์",
            "additional_source_screenshots": "ภาพต้นฉบับเพิ่มเติม",
            "appendix_note": "ภาพเหล่านี้มาจากเอกสารต้นฉบับ แต่ไม่ได้อยู่ในลำดับเนื้อหาหลัก",
            "appendix_images": "ภาพภาคผนวก",
        },
        "indonesian": {
            "introduction": "Pendahuluan",
            "conclusion": "Kesimpulan",
            "metadata_heading": "Metadata SEO",
            "metadata_field": "Kolom",
            "metadata_value": "Isi",
            "title_subtitle": "Dibuat dari paket artikel SEO dengan gambar sumber dipertahankan sesuai konteks.",
            "source_screenshots": "Gambar sumber",
            "source_caption_prefix": "Gambar sumber",
            "image_plan": "Rencana gambar",
            "image_label": "Gambar",
            "image_alt": "Teks alt",
            "image_purpose": "Tujuan",
            "additional_source_screenshots": "Gambar sumber tambahan",
            "appendix_note": "Gambar ini berasal dari dokumen sumber, tetapi tidak ditempatkan dalam alur artikel utama.",
            "appendix_images": "Gambar lampiran",
        },
        "vietnamese": {
            "introduction": "Giới thiệu",
            "conclusion": "Kết luận",
            "metadata_heading": "Thông tin SEO",
            "metadata_field": "Trường",
            "metadata_value": "Nội dung",
            "title_subtitle": "Được tạo từ gói bài viết SEO và giữ hình ảnh nguồn đúng theo ngữ cảnh.",
            "source_screenshots": "Hình ảnh nguồn",
            "source_caption_prefix": "Hình nguồn",
            "image_plan": "Kế hoạch hình ảnh",
            "image_label": "Hình ảnh",
            "image_alt": "Văn bản thay thế",
            "image_purpose": "Mục đích",
            "additional_source_screenshots": "Hình ảnh nguồn bổ sung",
            "appendix_note": "Các hình này đến từ tài liệu nguồn nhưng không được đặt trong luồng bài viết chính.",
            "appendix_images": "Hình ảnh phụ lục",
        },
    }
    community_packs = {
        "traditional_chinese": {
            "community_line_1": "想要更多福利、遊戲攻略和最新資訊？👉加入我們的 Discord！",
            "community_line_2": "超多實用內容、專屬福利只在社群內更新",
            "community_link_label": "🔗邀請連結：",
            "community_url": "https://discord.gg/Agkk96vcfA",
        },
        "english": {
            "community_line_1": "Want more rewards, game guides, and the latest updates? 👉 Join our Discord!",
            "community_line_2": "Useful tips and exclusive community benefits are updated only in our server.",
            "community_link_label": "🔗 Invitation link:",
            "community_url": "https://discord.gg/Agkk96vcfA",
        },
        "portuguese": {
            "community_line_1": "Quer mais benefícios, guias de jogos e as últimas novidades? 👉 Entre no nosso Discord!",
            "community_line_2": "Conteúdos úteis e benefícios exclusivos são atualizados somente na comunidade.",
            "community_link_label": "🔗 Link de convite:",
            "community_url": "https://discord.gg/FhSaQfq6rJ",
        },
        "spanish": {
            "community_line_1": "¿Quieres más beneficios, guías de juegos y las últimas novedades? 👉 ¡Únete a nuestro Discord!",
            "community_line_2": "Los mejores consejos y beneficios exclusivos se actualizan solo en la comunidad.",
            "community_link_label": "🔗 Enlace de invitación:",
            "community_url": "https://discord.gg/FhSaQfq6rJ",
        },
        "thai": {
            "community_line_1": "อยากรับสิทธิพิเศษ คู่มือเกม และข่าวสารล่าสุดเพิ่มเติมไหม? 👉 เข้าร่วม Discord ของเรา!",
            "community_line_2": "เคล็ดลับและสิทธิพิเศษเฉพาะจะอัปเดตภายในชุมชนเท่านั้น",
            "community_link_label": "🔗 ลิงก์เชิญ:",
            "community_url": "https://discord.gg/Agkk96vcfA",
        },
        "indonesian": {
            "community_line_1": "Ingin mendapatkan lebih banyak benefit, panduan game, dan info terbaru? 👉 Bergabunglah dengan Discord kami!",
            "community_line_2": "Tips bermanfaat dan benefit eksklusif hanya diperbarui di dalam komunitas.",
            "community_link_label": "🔗 Tautan undangan:",
            "community_url": "https://discord.gg/Agkk96vcfA",
        },
        "vietnamese": {
            "community_line_1": "Bạn muốn nhận thêm quyền lợi, hướng dẫn game và tin tức mới nhất? 👉 Tham gia Discord của chúng tôi!",
            "community_line_2": "Nhiều mẹo hữu ích và quyền lợi độc quyền chỉ được cập nhật trong cộng đồng.",
            "community_link_label": "🔗 Liên kết mời:",
            "community_url": "https://discord.gg/Agkk96vcfA",
        },
    }
    selected = dict(packs.get(language, packs["english"]))
    selected.update(community_packs.get(language, community_packs["english"]))
    return selected


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_indent(table, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        el = tc_mar.find(qn(f"w:{tag}"))
        if el is None:
            el = OxmlElement(f"w:{tag}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def style_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D9E2F2")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def parse_article_package(md_text: str):
    lines = md_text.splitlines()
    metadata = []
    image_plan = []
    article_title = ""
    article_sections = []
    current_section = None
    mode = None
    i = 0

    while i < len(lines):
        stripped = lines[i].strip()

        if stripped == "## SEO Metadata":
            mode = "metadata"
            i += 1
            continue
        if stripped == "## Article Body":
            mode = "article"
            i += 1
            continue
        if stripped == "## Image Plan":
            mode = "image_plan"
            i += 1
            continue

        if mode == "metadata":
            if stripped.startswith("- ") and ":" in stripped:
                key, value = stripped[2:].split(":", 1)
                metadata.append((key.strip(), value.strip().strip("`")))
            i += 1
            continue

        if mode == "article":
            if stripped.startswith("# "):
                article_title = stripped[2:].strip()
                i += 1
                continue
            if stripped.startswith("## "):
                if current_section:
                    article_sections.append(current_section)
                current_section = {"title": stripped[3:].strip(), "blocks": []}
                i += 1
                continue
            if stripped.startswith("### "):
                if current_section is None:
                    current_section = {"title": "__INTRO__", "blocks": []}
                current_section["blocks"].append(("h3", stripped[4:].strip()))
                i += 1
                continue
            image_marker = re.fullmatch(r"<!--\s*SOURCE_IMAGE\s*:\s*(\d+)\s*-->", stripped, re.IGNORECASE)
            if image_marker:
                if current_section is None:
                    current_section = {"title": "__INTRO__", "blocks": []}
                current_section["blocks"].append(("source_image", int(image_marker.group(1))))
                i += 1
                continue
            if re.match(r"^\d+\.\s+\*\*.+\*\*", stripped):
                if current_section is None:
                    current_section = {"title": "__INTRO__", "blocks": []}
                lead = re.sub(r"^\d+\.\s+\*\*(.+?)\*\*\s*$", r"\1", stripped)
                item_lines = [lead]
                j = i + 1
                while j < len(lines) and lines[j].startswith("   "):
                    item_lines.append(lines[j].strip())
                    j += 1
                current_section["blocks"].append(("number_item", " ".join(item_lines)))
                i = j
                continue
            if stripped.startswith("- "):
                if current_section is None:
                    current_section = {"title": "__INTRO__", "blocks": []}
                items = []
                while i < len(lines) and lines[i].strip().startswith("- "):
                    items.append(lines[i].strip()[2:].strip())
                    i += 1
                current_section["blocks"].append(("bullets", items))
                continue
            if stripped:
                if current_section is None:
                    current_section = {"title": "__INTRO__", "blocks": []}
                para_lines = [stripped]
                j = i + 1
                while j < len(lines):
                    nxt = lines[j].strip()
                    if not nxt:
                        break
                    if nxt.startswith("#") or nxt.startswith("## ") or nxt.startswith("### ") or nxt.startswith("- "):
                        break
                    if re.fullmatch(r"<!--\s*SOURCE_IMAGE\s*:\s*\d+\s*-->", nxt, re.IGNORECASE):
                        break
                    if re.match(r"^\d+\.\s+\*\*.+", nxt):
                        break
                    para_lines.append(nxt)
                    j += 1
                current_section["blocks"].append(("paragraph", " ".join(para_lines)))
                i = j
                continue
            i += 1
            continue

        if mode == "image_plan":
            if re.match(r"^\d+\.\s+", stripped):
                title = re.sub(r"^\d+\.\s+", "", stripped)
                alt = ""
                purpose = ""
                source_image = None
                j = i + 1
                while j < len(lines) and lines[j].startswith("   - "):
                    detail = lines[j].strip()[2:].strip()
                    if detail.startswith("Alt:"):
                        alt = detail.split(":", 1)[1].strip().strip("`")
                    elif detail.startswith("Purpose:") or detail.startswith("用途:") or detail.startswith("Uso:"):
                        purpose = detail.split(":", 1)[1].strip()
                    elif any(detail.startswith(prefix) for prefix in ("Source Image:", "原图序号:", "Imagen de origen:")):
                        match = re.search(r"\d+", detail)
                        source_image = int(match.group(0)) if match else None
                    j += 1
                image_plan.append({
                    "title": title,
                    "alt": alt,
                    "purpose": purpose,
                    "source_image": source_image,
                })
                i = j
                continue
            i += 1
            continue

        i += 1

    if current_section:
        article_sections.append(current_section)

    return article_title, metadata, article_sections, image_plan


def extract_source_images(source_docx: Path):
    source_doc = Document(source_docx)
    images = []
    for index, shape in enumerate(source_doc.inline_shapes, 1):
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        part = source_doc.part.related_parts[blip.embed]
        images.append({"index": index, "name": Path(part.partname).name, "blob": part.blob})
    return images


def verify_built_docx(output_path: Path, expected_images, labels: dict):
    if not output_path.exists() or output_path.stat().st_size == 0:
        raise ValueError(f"Generated DOCX is missing or empty: {output_path}")

    try:
        built_doc = Document(output_path)
    except Exception as exc:
        raise ValueError(f"Generated DOCX cannot be reopened: {output_path}") from exc

    if not any(paragraph.text.strip() for paragraph in built_doc.paragraphs):
        raise ValueError("Generated DOCX contains no readable paragraph text.")

    document_text = "\n".join(paragraph.text for paragraph in built_doc.paragraphs)
    for required_text in (
        labels["community_line_1"],
        labels["community_line_2"],
        labels["community_link_label"],
        labels["community_url"],
    ):
        if required_text not in document_text:
            raise ValueError(f"Generated DOCX is missing community invitation content: {required_text}")

    hyperlink_targets = {
        rel.target_ref
        for rel in built_doc.part.rels.values()
        if rel.reltype == RT.HYPERLINK and rel.is_external
    }
    if labels["community_url"] not in hyperlink_targets:
        raise ValueError("Generated DOCX community invitation URL is not a clickable external hyperlink.")

    built_images = extract_source_images(output_path)
    if len(built_images) != len(expected_images):
        raise ValueError(
            "Generated DOCX image count does not match the source image count. "
            f"Expected {len(expected_images)}, found {len(built_images)}."
        )

    expected_hashes = [hashlib.sha256(image["blob"]).hexdigest() for image in expected_images]
    built_hashes = [hashlib.sha256(image["blob"]).hexdigest() for image in built_images]
    if built_hashes != expected_hashes:
        raise ValueError("Generated DOCX source images are missing, altered, or out of order.")


def build_doc_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "#2E74B5", 18, 10),
        ("Heading 2", 13, "#2E74B5", 14, 7),
        ("Heading 3", 12, "#1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.25


def add_title_block(doc: Document, title: str, labels: dict):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(para, before=0, after=10, line=1.15)
    run = para.add_run(title)
    set_run_font(run, size=20, bold=True, color="#2E74B5")

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, before=0, after=10, line=1.15)
    run = subtitle.add_run(labels["title_subtitle"])
    set_run_font(run, size=10, color="#5B6573")


def add_hyperlink(paragraph, text: str, url: str):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_community_invitation(doc: Document, labels: dict):
    first = doc.add_paragraph()
    set_paragraph_spacing(first, before=14, after=4, line=1.25)
    set_run_font(first.add_run(labels["community_line_1"]), bold=True)

    second = doc.add_paragraph()
    set_paragraph_spacing(second, before=0, after=4, line=1.25)
    set_run_font(second.add_run(labels["community_line_2"]))

    link_paragraph = doc.add_paragraph()
    set_paragraph_spacing(link_paragraph, before=0, after=6, line=1.25)
    set_run_font(link_paragraph.add_run(f'{labels["community_link_label"]} '))
    add_hyperlink(link_paragraph, labels["community_url"], labels["community_url"])


def add_metadata_table(doc: Document, metadata, labels: dict):
    heading = doc.add_paragraph(labels["metadata_heading"], style="Heading 2")
    set_paragraph_spacing(heading, before=14, after=7, line=1.25)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_indent(table, 120)
    style_table_borders(table)
    widths = [2700, 6660]

    for idx, text in enumerate([labels["metadata_field"], labels["metadata_value"]]):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        shade_cell(cell, "E8EEF5")
        para = cell.paragraphs[0]
        set_paragraph_spacing(para, before=0, after=0, line=1.15)
        set_run_font(para.add_run(text), bold=True)

    for key, value in metadata:
        row = table.add_row().cells
        for idx, cell in enumerate(row):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
        p0 = row[0].paragraphs[0]
        set_paragraph_spacing(p0, before=0, after=0, line=1.15)
        set_run_font(p0.add_run(key), bold=True)
        p1 = row[1].paragraphs[0]
        set_paragraph_spacing(p1, before=0, after=0, line=1.15)
        set_run_font(p1.add_run(value))


def add_section_blocks(doc: Document, section, labels: dict, source_images_by_index=None, image_alt_by_index=None):
    if section["title"] == "__INTRO__":
        heading_text = labels["introduction"]
    elif re.fullmatch(
        r"(?:Conclusion(?:\s*\+\s*CTA)?|Conclusão|Conclusión|结语(?:与\s*CTA)?|結語(?:與\s*CTA)?|结论(?:与\s*CTA)?|結論(?:與\s*CTA)?|สรุป|Kesimpulan|Kết luận)",
        section["title"],
        re.IGNORECASE,
    ):
        heading_text = labels["conclusion"]
    else:
        heading_text = section["title"]
    heading = doc.add_paragraph(heading_text, style="Heading 1")
    set_paragraph_spacing(heading, before=18, after=10, line=1.25)

    pending_numbers = []
    for block_type, payload in section["blocks"]:
        if block_type == "number_item":
            pending_numbers.append(payload)
            continue

        if pending_numbers:
            for item in pending_numbers:
                para = doc.add_paragraph(style="List Number")
                set_paragraph_spacing(para, before=0, after=4, line=1.25)
                set_run_font(para.add_run(item))
            pending_numbers = []

        if block_type == "paragraph":
            para = doc.add_paragraph()
            set_paragraph_spacing(para, before=0, after=6, line=1.25)
            set_run_font(para.add_run(payload))
        elif block_type == "h3":
            para = doc.add_paragraph(payload, style="Heading 3")
            set_paragraph_spacing(para, before=10, after=5, line=1.25)
        elif block_type == "bullets":
            for item in payload:
                para = doc.add_paragraph(style="List Bullet")
                set_paragraph_spacing(para, before=0, after=4, line=1.25)
                set_run_font(para.add_run(item))
        elif block_type == "source_image":
            image = (source_images_by_index or {}).get(payload)
            if image is None:
                raise ValueError(f"SOURCE_IMAGE:{payload} does not exist in the source document.")
            alt_text = (image_alt_by_index or {}).get(payload, "")
            caption = alt_text or f"{labels['source_caption_prefix']} {payload:02d} - {image['name']}"
            add_picture_with_caption(doc, image["blob"], caption, alt_text=alt_text)

    if pending_numbers:
        for item in pending_numbers:
            para = doc.add_paragraph(style="List Number")
            set_paragraph_spacing(para, before=0, after=4, line=1.25)
            set_run_font(para.add_run(item))


def add_picture_with_caption(doc: Document, image_blob: bytes, caption: str, width=Inches(6.15), alt_text=""):
    pic_para = doc.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(pic_para, before=4, after=3, line=1.0)
    inline_shape = pic_para.add_run().add_picture(BytesIO(image_blob), width=width)
    if alt_text:
        doc_pr = inline_shape._inline.docPr
        doc_pr.set("descr", alt_text)
        doc_pr.set("title", alt_text)

    cap_para = doc.add_paragraph()
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cap_para, before=0, after=8, line=1.0)
    set_run_font(cap_para.add_run(caption), size=9, color="#5B6573", italic=True)


def collect_source_image_markers(article_sections):
    return [
        payload
        for section in article_sections
        for block_type, payload in section["blocks"]
        if block_type == "source_image"
    ]


def validate_source_image_markers(source_images, article_sections):
    markers = collect_source_image_markers(article_sections)
    expected = list(range(1, len(source_images) + 1))
    if markers != expected:
        raise ValueError(
            "Source image placement must preserve every image exactly once and in order. "
            f"Expected markers {expected}, found {markers}. Add standalone markers such as "
            "<!-- SOURCE_IMAGE:1 --> inside the matching rewritten sections."
        )


def add_image_gallery(doc: Document, heading_text: str, images, labels: dict):
    if not images:
        return
    gallery_heading = doc.add_paragraph(heading_text, style="Heading 3")
    set_paragraph_spacing(gallery_heading, before=10, after=6, line=1.25)
    for image in images:
        add_picture_with_caption(
            doc,
            image["blob"],
            f"{labels['source_caption_prefix']} {image['index']:02d} - {image['name']}",
        )


def add_image_plan_table(doc: Document, image_plan, labels: dict):
    heading = doc.add_paragraph(labels["image_plan"], style="Heading 1")
    set_paragraph_spacing(heading, before=18, after=10, line=1.25)

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    set_table_layout_fixed(table)
    set_table_indent(table, 120)
    style_table_borders(table)
    widths = [2400, 2900, 4060]

    for idx, text in enumerate([labels["image_label"], labels["image_alt"], labels["image_purpose"]]):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        shade_cell(cell, "E8EEF5")
        para = cell.paragraphs[0]
        set_paragraph_spacing(para, before=0, after=0, line=1.15)
        set_run_font(para.add_run(text), bold=True)

    for item in image_plan:
        title = item["title"]
        alt = item["alt"]
        purpose = item["purpose"]
        row = table.add_row().cells
        for idx, value in enumerate([title, alt, purpose]):
            cell = row[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            set_paragraph_spacing(para, before=0, after=0, line=1.15)
            set_run_font(para.add_run(value))


def build_docx(article_package: Path, output_path: Path, source_docx: Path | None):
    article_title, metadata, article_sections, image_plan = parse_article_package(article_package.read_text(encoding="utf-8"))
    metadata_map = metadata_to_dict(metadata)
    labels = get_language_pack(metadata_map.get("Output Language", "english"))
    source_images = extract_source_images(source_docx) if source_docx else []
    if source_images:
        validate_source_image_markers(source_images, article_sections)
    elif collect_source_image_markers(article_sections):
        raise ValueError("The article package contains SOURCE_IMAGE markers, but no source document images were provided.")
    source_images_by_index = {image["index"]: image for image in source_images}
    image_alt_by_index = {
        item["source_image"]: item["alt"]
        for item in image_plan
        if item["source_image"] is not None and item["alt"]
    }

    doc = Document()
    build_doc_styles(doc)
    add_title_block(doc, article_title, labels)
    add_metadata_table(doc, metadata, labels)

    for section in article_sections:
        add_section_blocks(doc, section, labels, source_images_by_index, image_alt_by_index)

    add_community_invitation(doc, labels)

    doc.save(output_path)
    verify_built_docx(output_path, source_images, labels)


def main():
    parser = argparse.ArgumentParser(description="Build a Word document from a game-guides article package.")
    parser.add_argument("--article-package", required=True, type=Path, help="Markdown article package with Title, SEO Metadata, Article Body, and Image Plan.")
    parser.add_argument("--out", required=True, type=Path, help="Output .docx path.")
    parser.add_argument("--source-docx", type=Path, help="Optional source .docx to preserve screenshots from.")
    args = parser.parse_args()

    build_docx(args.article_package, args.out, args.source_docx)
    print(args.out)


if __name__ == "__main__":
    main()
