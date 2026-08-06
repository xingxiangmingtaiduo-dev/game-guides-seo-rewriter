from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def configure_console() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


def iter_paragraph_tokens(paragraph, document, image_state):
    text_parts = []
    for node in paragraph._p.iter():
        if node.tag == qn("w:t") and node.text:
            text_parts.append(node.text)
        elif node.tag == qn("w:tab"):
            text_parts.append("\t")
        elif node.tag == qn("w:br"):
            text_parts.append("\n")
        elif node.tag == qn("a:blip"):
            text = "".join(text_parts).strip()
            if text:
                yield ("text", text)
            text_parts = []
            relationship_id = node.get(qn("r:embed"))
            if not relationship_id:
                continue
            image_state["index"] += 1
            part = document.part.related_parts[relationship_id]
            yield (
                "image",
                {
                    "index": image_state["index"],
                    "name": Path(part.partname).name,
                },
            )
    text = "".join(text_parts).strip()
    if text:
        yield ("text", text)


def inspect_document(source_docx: Path) -> str:
    document = Document(source_docx)
    image_state = {"index": 0}
    output = [f"# Source sequence: {source_docx.name}", ""]

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = next((p for p in document.paragraphs if p._p is child), None)
            if paragraph is None:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            for token_type, payload in iter_paragraph_tokens(paragraph, document, image_state):
                if token_type == "image":
                    output.append(f"<!-- SOURCE_IMAGE:{payload['index']} -->")
                    output.append(f"[Original file: {payload['name']}]")
                elif payload:
                    prefix = "## " if style_name.startswith("Heading") else ""
                    output.append(f"{prefix}{payload}")
                output.append("")
        elif child.tag == qn("w:tbl"):
            table = next((t for t in document.tables if t._tbl is child), None)
            if table is None:
                continue
            output.append("[Table]")
            output.append("")
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_parts = []
                    for paragraph in cell.paragraphs:
                        for token_type, payload in iter_paragraph_tokens(paragraph, document, image_state):
                            if token_type == "image":
                                if cell_parts:
                                    row_text.append(" ".join(cell_parts))
                                    cell_parts = []
                                row_text.append(f"<!-- SOURCE_IMAGE:{payload['index']} -->")
                            elif payload:
                                cell_parts.append(payload)
                    if cell_parts:
                        row_text.append(" ".join(cell_parts))
                if row_text:
                    output.append(" | ".join(row_text))
                    output.append("")

    output.append(f"Total source images: {image_state['index']}")
    return "\n".join(output).strip() + "\n"


def main() -> None:
    configure_console()
    parser = argparse.ArgumentParser(
        description="Print source DOCX text and image markers in their original reading order."
    )
    parser.add_argument("--source-docx", required=True, type=Path)
    parser.add_argument("--out", type=Path, help="Optional UTF-8 Markdown output path.")
    args = parser.parse_args()

    report = inspect_document(args.source_docx)
    if args.out:
        args.out.write_text(report, encoding="utf-8")
        print(args.out)
    else:
        print(report, end="")


if __name__ == "__main__":
    main()
