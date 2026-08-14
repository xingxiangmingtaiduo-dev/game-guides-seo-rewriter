from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from docx import Document

from build_docx import build_docx, get_language_pack, normalize_output_language
from build_markdown import build_markdown
from validate_article_package import count_words, normalize_output_language as normalize_validator_language
from validate_article_package import parse_article_package, validate


LANGUAGES = {
    "Traditional Chinese": ("traditional_chinese", "前言", "結語", "SEO 資訊"),
    "English": ("english", "Introduction", "Conclusion", "SEO Metadata"),
    "Portuguese": ("portuguese", "Introdução", "Conclusão", "Metadados SEO"),
    "Spanish": ("spanish", "Introducción", "Conclusión", "Metadatos SEO"),
    "Thai": ("thai", "บทนำ", "สรุป", "ข้อมูล SEO"),
    "Indonesian": ("indonesian", "Pendahuluan", "Kesimpulan", "Metadata SEO"),
    "Vietnamese": ("vietnamese", "Giới thiệu", "Kết luận", "Thông tin SEO"),
}

ALIASES = {
    "繁體中文": "traditional_chinese",
    "English": "english",
    "Português": "portuguese",
    "Español": "spanish",
    "ภาษาไทย": "thai",
    "Bahasa Indonesia": "indonesian",
    "Tiếng Việt": "vietnamese",
}

VALIDATION_TEXT = {
    "Traditional Chinese": ("Test是什麼遊戲？", "如何使用 UgPhone？", "結語", "立即下載UgPhone並申請免費試用"),
    "English": ("What is Test?", "How to Use UgPhone?", "Conclusion", "download UgPhone and request a free trial"),
    "Portuguese": ("O que é Test?", "Como usar UgPhone?", "Conclusão", "baixar UgPhone e solicitar um teste gratuito"),
    "Spanish": ("¿Qué es Test?", "¿Cómo usar UgPhone?", "Conclusión", "descargar UgPhone y solicitar una prueba gratuita"),
    "Thai": ("Test คือเกมอะไร?", "ใช้ UgPhone อย่างไร?", "สรุป", "ดาวน์โหลด UgPhone และขอทดลองใช้ฟรี"),
    "Indonesian": ("Apa itu Test?", "Cara menggunakan UgPhone?", "Kesimpulan", "unduh UgPhone dan minta uji coba gratis"),
    "Vietnamese": ("Test là gì?", "Cách dùng UgPhone?", "Kết luận", "tải UgPhone và dùng thử miễn phí"),
}


def article_package(language: str) -> str:
    what_heading, ugphone_heading, conclusion_heading, cta = VALIDATION_TEXT[language]
    return f"""# Title

Multilingual smoke test

## SEO Metadata

- Slug: multilingual-smoke-test
- SEO Title: Multilingual smoke test
- Primary Keyword: smoke test
- Meta Description: This metadata exists to verify localized document labels and multilingual output without exposing internal image planning data.
- Tags: test, UgPhone
- Output Language: {language}

## Article Body

# Multilingual smoke test

Opening paragraph for the localized introduction heading.

## {what_heading}

It verifies the public document structure.

## {ugphone_heading}

Run a short session and check the result.

## {conclusion_heading}

{cta} after verification.

## Image Plan

1. Test image one
   - Alt: smoke test example one
   - Purpose: Verify parsing.
2. Test image two
   - Alt: test example two
   - Purpose: Verify parsing.
3. Test image three
   - Alt: test example three
   - Purpose: Verify parsing.
4. Test image four
   - Alt: test example four
   - Purpose: Verify parsing.
5. Test image five
   - Alt: test example five
   - Purpose: Verify parsing.
6. Test image six
   - Alt: test example six
   - Purpose: Verify parsing.
"""


def main() -> None:
    vietnamese_sample = "Hướng dẫn tiếng Việt giúp người chơi kiểm tra kết nối và phần thưởng."
    assert count_words(vietnamese_sample) == 14, count_words(vietnamese_sample)
    for alias, normalized in ALIASES.items():
        assert normalize_output_language(alias) == normalized
        assert normalize_validator_language(alias) == normalized

    with tempfile.TemporaryDirectory(prefix="seo-language-smoke-") as temp:
        root = Path(temp)
        for language, (normalized, introduction, conclusion, metadata_heading) in LANGUAGES.items():
            assert normalize_output_language(language) == normalized
            assert normalize_validator_language(language) == normalized
            labels = get_language_pack(language)
            assert labels["introduction"] == introduction
            assert labels["conclusion"] == conclusion
            assert labels["metadata_heading"] == metadata_heading
            expected_community_url = (
                "https://discord.gg/FhSaQfq6rJ"
                if normalized in {"portuguese", "spanish"}
                else "https://discord.gg/Agkk96vcfA"
            )
            wrong_community_url = (
                "https://discord.gg/Agkk96vcfA"
                if normalized in {"portuguese", "spanish"}
                else "https://discord.gg/FhSaQfq6rJ"
            )
            assert labels["community_url"] == expected_community_url

            package = root / f"{normalized}.md"
            docx_path = root / f"{normalized}.docx"
            markdown_path = root / f"{normalized}-public.md"
            package.write_text(article_package(language), encoding="utf-8")
            checks = validate(parse_article_package(package.read_text(encoding="utf-8")), "game-guide")
            required = {
                "output_language",
                "conclusion_heading_without_cta",
                "core_topic_section",
                "main_action_section",
                "ugphone_value_section",
                "ugphone_tutorial_or_flow",
                "cta_present",
            }
            failed = [check for check in checks if check["name"] in required and not check["passed"]]
            assert not failed, (language, failed)
            build_docx(package, docx_path, None)
            build_markdown(package, markdown_path)

            docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
            markdown_text = markdown_path.read_text(encoding="utf-8")
            assert introduction in docx_text and conclusion in docx_text and metadata_heading in docx_text
            assert f"## {introduction}" in markdown_text and f"## {conclusion}" in markdown_text
            assert "Image Plan" not in docx_text and "Image Plan" not in markdown_text
            assert labels["community_line_1"] in docx_text and labels["community_line_2"] in docx_text
            assert labels["community_line_1"] in markdown_text and labels["community_line_2"] in markdown_text
            assert expected_community_url in docx_text and wrong_community_url not in docx_text
            assert f"[{expected_community_url}]({expected_community_url})" in markdown_text
            assert wrong_community_url not in markdown_text

            built_doc = Document(docx_path)
            hyperlink_targets = {
                rel.target_ref
                for rel in built_doc.part.rels.values()
                if rel.is_external and rel.reltype.endswith("/hyperlink")
            }
            assert expected_community_url in hyperlink_targets

    print("PASS: 7 multilingual DOCX/Markdown language packs with localized Discord invitations")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        raise
