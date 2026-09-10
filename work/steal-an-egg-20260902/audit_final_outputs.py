import hashlib
import re
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfReader


OUT = Path(r"C:\Users\Og\Desktop\已SEO文章\Steal an egg七语言SEO")
SOURCE = Path(r"C:\Users\Og\Desktop\待SEO文章\Steal an egg测评.docx")

LANGUAGES = {
    "Traditional-Chinese": {
        "slug": "steal-an-egg-afk-guide-tw",
        "title": "Steal an egg掛機攻略：偷蛋、孵蛋與練速",
        "intro": "前言",
        "conclusion": "結語",
        "metadata": "SEO 資訊",
        "url": "https://discord.gg/Agkk96vcfA",
    },
    "English": {
        "slug": "steal-an-egg-afk-guide-en",
        "title": "Steal an egg AFK Guide: Eggs, Pets, and Speed",
        "intro": "Introduction",
        "conclusion": "Conclusion",
        "metadata": "SEO Metadata",
        "url": "https://discord.gg/Agkk96vcfA",
    },
    "Portuguese": {
        "slug": "steal-an-egg-afk-guide-pt",
        "title": "Guia AFK de Steal an egg: ovos, pets e velocidade",
        "intro": "Introdução",
        "conclusion": "Conclusão",
        "metadata": "Metadados SEO",
        "url": "https://discord.gg/FhSaQfq6rJ",
    },
    "Spanish": {
        "slug": "steal-an-egg-afk-guide-es",
        "title": "Guía AFK de Steal an egg: huevos, mascotas y velocidad",
        "intro": "Introducción",
        "conclusion": "Conclusión",
        "metadata": "Metadatos SEO",
        "url": "https://discord.gg/FhSaQfq6rJ",
    },
    "Thai": {
        "slug": "steal-an-egg-afk-guide-th",
        "title": "คู่มือ AFK Steal an egg: ไข่ สัตว์เลี้ยง และความเร็ว",
        "intro": "บทนำ",
        "conclusion": "สรุป",
        "metadata": "ข้อมูล SEO",
        "url": "https://discord.gg/Agkk96vcfA",
    },
    "Indonesian": {
        "slug": "steal-an-egg-afk-guide-id",
        "title": "Panduan AFK Steal an egg: telur, pet, dan kecepatan",
        "intro": "Pendahuluan",
        "conclusion": "Kesimpulan",
        "metadata": "Metadata SEO",
        "url": "https://discord.gg/Agkk96vcfA",
    },
    "Vietnamese": {
        "slug": "steal-an-egg-afk-guide-vi",
        "title": "Hướng dẫn AFK Steal an egg: trứng, pet và tốc độ",
        "intro": "Giới thiệu",
        "conclusion": "Kết luận",
        "metadata": "Thông tin SEO",
        "url": "https://discord.gg/Agkk96vcfA",
    },
}


def image_hashes(path: Path) -> list[str]:
    doc = Document(path)
    hashes = []
    for shape in doc.inline_shapes:
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        hashes.append(hashlib.sha256(doc.part.related_parts[blip.embed].blob).hexdigest())
    return hashes


expected_files = {
    f"Steal_an_egg_SEO_{language}_500-600{suffix}"
    for language in LANGUAGES
    for suffix in (".docx", ".md")
}
actual_files = {path.name for path in OUT.iterdir() if path.is_file()}
assert actual_files == expected_files, (actual_files - expected_files, expected_files - actual_files)

expected_dirs = {
    f"Steal_an_egg_SEO_{language}_500-600_assets" for language in LANGUAGES
}
actual_dirs = {path.name for path in OUT.iterdir() if path.is_dir()}
assert actual_dirs == expected_dirs, (actual_dirs - expected_dirs, expected_dirs - actual_dirs)

source_hashes = image_hashes(SOURCE)
assert len(source_hashes) == 15

for language, expected in LANGUAGES.items():
    stem = f"Steal_an_egg_SEO_{language}_500-600"
    docx_path = OUT / f"{stem}.docx"
    md_path = OUT / f"{stem}.md"
    assets_path = OUT / f"{stem}_assets"

    doc = Document(docx_path)
    doc_text_parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        doc_text_parts.extend(cell.text for row in table.rows for cell in row.cells)
    doc_text = "\n".join(doc_text_parts)
    assert "�" not in doc_text
    assert "SOURCE_IMAGE" not in doc_text and "Image Plan" not in doc_text
    for value in (expected["slug"], expected["title"], expected["intro"], expected["conclusion"], expected["metadata"], expected["url"]):
        assert value in doc_text, (language, value)
    hyperlinks = {
        rel.target_ref
        for rel in doc.part.rels.values()
        if rel.reltype == RT.HYPERLINK and rel.is_external
    }
    assert expected["url"] in hyperlinks
    wrong_url = "https://discord.gg/Agkk96vcfA" if expected["url"].endswith("FhSaQfq6rJ") else "https://discord.gg/FhSaQfq6rJ"
    assert wrong_url not in doc_text and wrong_url not in hyperlinks
    assert image_hashes(docx_path) == source_hashes

    md_bytes = md_path.read_bytes()
    md_text = md_bytes.decode("utf-8", errors="strict")
    assert "�" not in md_text
    assert "SOURCE_IMAGE" not in md_text and "## Image Plan" not in md_text
    assert f'slug: "{expected["slug"]}"' in md_text
    assert f'# {expected["title"]}' in md_text
    assert f'## {expected["intro"]}' in md_text
    assert f'## {expected["conclusion"]}' in md_text
    assert f'[{expected["url"]}]({expected["url"]})' in md_text
    assert wrong_url not in md_text
    links = re.findall(r"^!\[.*?\]\(<(.+?)>\)\s*$", md_text, re.MULTILINE)
    assert len(links) == 15
    asset_files = sorted(assets_path.glob("source-image-*"))
    assert len(asset_files) == 15
    assert [hashlib.sha256(path.read_bytes()).hexdigest() for path in asset_files] == source_hashes

    if language not in {"Traditional-Chinese"}:
        assert not re.search(r"[\u3400-\u9fff]", md_text), f"Unexpected CJK text in {language}"

    print(f"PASS {language}: 15 images, slug {expected['slug']}, localized metadata and CTA")

print("PASS ALL: 7 DOCX, 7 Markdown, 7 asset directories")

pdf_dir = Path(__file__).parent / "qa-pdf"
pdf_files = sorted(pdf_dir.glob("*.pdf"))
assert len(pdf_files) == 7
for pdf_path in pdf_files:
    reader = PdfReader(pdf_path)
    extracted = [page.extract_text() or "" for page in reader.pages]
    assert reader.pages and all(text.strip() for text in extracted)
    assert sum(len(text) for text in extracted) > 500
    print(
        f"PASS PDF {pdf_path.stem}: pages={len(reader.pages)}, "
        f"text_chars={sum(len(text) for text in extracted)}"
    )
