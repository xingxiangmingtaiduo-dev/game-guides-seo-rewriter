from pathlib import Path

from docx import Document


OUT = Path(r"C:\Users\Og\Desktop\the_tower_seo_special_20260928")


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def replace_first_heading(doc, old, new):
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip() == old:
            replace_paragraph(paragraph, new)
            return
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip() == new:
            return
    raise ValueError(f"Heading not found: {old}")


def patch_metadata(doc, language):
    if language == "de":
        subtitle = "Aus dem SEO-Artikelpaket erstellt. Quell-Screenshots bleiben an ihrer passenden Stelle im Artikel erhalten."
        heading = "SEO-Metadaten"
        headers = ("Feld", "Wert")
        labels = {
            "SEO Title": "SEO-Titel",
            "Primary Keyword": "Hauptkeyword",
            "Long-tail Keywords": "Long-Tail-Keywords",
            "Meta Description": "Meta-Beschreibung",
            "Body Length Target": "Zieltextlänge",
            "Output Language": "Ausgabesprache",
        }
        values = {
            "Slug": "the-tower-afk-guide-de",
            "Long-tail Keywords": "The Tower AFK Guide AFK, The Tower-Münzen, The Tower-Edelsteine, UgPhone The Tower",
            "Tags": "The Tower, AFK-Guide, Münzen, Edelsteine, UgPhone",
            "Output Language": "Deutsch",
        }
    else:
        subtitle = "Généré à partir du pack SEO de l’article. Les captures d’écran sources restent à leur emplacement dans l’article."
        heading = "Métadonnées SEO"
        headers = ("Champ", "Valeur")
        labels = {
            "SEO Title": "Titre SEO",
            "Primary Keyword": "Mot-clé principal",
            "Long-tail Keywords": "Mots-clés longue traîne",
            "Meta Description": "Méta-description",
            "Body Length Target": "Longueur cible",
            "Output Language": "Langue de sortie",
        }
        values = {
            "Slug": "the-tower-afk-guide-fr",
            "Long-tail Keywords": "guide AFK The Tower AFK, pièces The Tower, gemmes The Tower, UgPhone The Tower",
            "Tags": "The Tower, guide AFK, pièces, gemmes, UgPhone",
            "Output Language": "Français",
        }

    for paragraph in doc.paragraphs[:4]:
        if paragraph.style.name == "Heading 2" and paragraph.text.strip() == "SEO Metadata":
            replace_paragraph(paragraph, heading)
        elif paragraph.text.strip().startswith("Generated from the SEO article package"):
            replace_paragraph(paragraph, subtitle)

    if not doc.tables:
        return
    table = doc.tables[0]
    replace_paragraph(table.rows[0].cells[0].paragraphs[0], headers[0])
    replace_paragraph(table.rows[0].cells[1].paragraphs[0], headers[1])
    for row in table.rows[1:]:
        key = row.cells[0].text.strip()
        if key in labels:
            replace_paragraph(row.cells[0].paragraphs[0], labels[key])
        if key in values:
            replace_paragraph(row.cells[1].paragraphs[0], values[key])


def patch_docx(path, language, intro_heading, conclusion_heading, line1, line2, link_label):
    doc = Document(path)
    replace_first_heading(doc, "Introduction", intro_heading)
    patch_metadata(doc, language)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "Conclusion":
            replace_paragraph(paragraph, conclusion_heading)
    replace_paragraph(doc.paragraphs[-3], line1)
    replace_paragraph(doc.paragraphs[-2], line2)
    if doc.paragraphs[-1].runs:
        doc.paragraphs[-1].runs[0].text = link_label + " "
    doc.save(path)


patch_docx(
    OUT / "de.docx", "de", "Einleitung", "Fazit",
    "Möchtest du mehr Belohnungen, Spielguides und aktuelle Neuigkeiten? 👉 Tritt unserem Discord bei!",
    "Nützliche Tipps und exklusive Vorteile werden nur in unserer Community aktualisiert.",
    "🔗 Einladungslink:",
)
patch_docx(
    OUT / "fr.docx", "fr", "Présentation", "Conclusion",
    "Vous voulez plus de récompenses, de guides de jeux et les dernières nouveautés ? 👉 Rejoignez notre Discord !",
    "Des conseils utiles et des avantages exclusifs sont mis à jour uniquement dans notre communauté.",
    "🔗 Lien d’invitation :",
)
print("fixed")
